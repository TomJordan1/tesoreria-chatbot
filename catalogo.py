"""Configuración de proyectos sin inferencias de IA ni selección por calendario."""
from dataclasses import dataclass
from io import BytesIO
import math
from urllib.parse import quote


class ErrorCatalogo(ValueError):
    """El catálogo no permite determinar un destino seguro."""


def texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and math.isfinite(valor) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def activo(valor):
    return texto(valor).lower() in ("sí", "si", "true", "1")


def identificador(valor):
    valor = texto(valor)
    if not valor or any(c in valor for c in "/\\?#"):
        raise ErrorCatalogo("Usa identificadores o nombres de tabla, no URLs ni rutas.")
    return valor


def telegram_id(valor):
    valor = texto(valor)
    if not valor.isascii() or not valor.isdigit() or int(valor) <= 0:
        raise ErrorCatalogo("telegram_id debe ser un entero positivo escrito como texto.")
    return str(int(valor))


def leer_tablas_xlsx(contenido):
    """Lee tablas reales de un archivo descargado; nunca guarda ni ejecuta fórmulas."""
    from openpyxl import load_workbook

    libro = load_workbook(BytesIO(contenido), data_only=False, keep_links=False)
    tablas = {}
    try:
        for hoja in libro:
            for nombre in hoja.tables:
                tabla = hoja.tables[nombre]
                filas = list(hoja[tabla.ref])
                if nombre in tablas:
                    raise ErrorCatalogo("Los nombres de las tablas deben ser únicos.")
                if any(celda.data_type == "f" for fila in filas for celda in fila):
                    raise ErrorCatalogo("La configuración debe contener valores, no fórmulas.")
                valores = [[celda.value for celda in fila] for fila in filas]
                if tabla.totalsRowCount:
                    valores = valores[:-tabla.totalsRowCount]
                tablas[nombre] = valores
        return tablas
    finally:
        libro.close()


def registros(tablas, nombre, campos):
    filas = tablas.get(nombre)
    if not filas or not isinstance(filas[0], list):
        raise ErrorCatalogo(f"Falta la tabla {nombre} con sus encabezados.")
    encabezados = [texto(v).lower() for v in filas[0]]
    if len(set(encabezados)) != len(encabezados) or not set(campos) <= set(encabezados):
        raise ErrorCatalogo(f"Encabezados incorrectos en {nombre}: {', '.join(campos)}.")
    for fila in filas[1:]:
        if not any(texto(v) for v in fila):
            continue
        if len(fila) != len(encabezados):
            raise ErrorCatalogo(f"Fila incompleta en {nombre}.")
        yield dict(zip(encabezados, fila))


@dataclass(frozen=True)
class Destino:
    proyecto_id: str
    nombre: str
    gestion: str
    drive_id: str
    item_id: str
    hoja: str
    tabla: str
    formato_fecha: str

    def url_tabla(self):
        partes = [quote(v, safe="") for v in (self.drive_id, self.item_id, self.tabla)]
        return f"https://graph.microsoft.com/v1.0/drives/{partes[0]}/items/{partes[1]}/workbook/tables/{partes[2]}"


class Catalogo:
    def __init__(self, tablas, archivo_config=None):
        self.usuarios = {}
        for fila in registros(tablas, "USUARIOS_AUTORIZADOS", ("telegram_id", "nombre", "activo")):
            if not texto(fila["telegram_id"]) and not activo(fila["activo"]):
                continue
            uid = telegram_id(fila["telegram_id"])
            self.usuarios[uid] = self.usuarios.get(uid, True) and activo(fila["activo"])

        self.proyectos = {}
        for fila in registros(tablas, "PROYECTOS", ("proyecto_id", "nombre", "gestion_actual", "activo")):
            pid = identificador(fila["proyecto_id"])
            if pid in self.proyectos:
                raise ErrorCatalogo(f"Proyecto duplicado: {pid}.")
            self.proyectos[pid] = (texto(fila["nombre"]), texto(fila["gestion_actual"]), activo(fila["activo"]))
            if not self.proyectos[pid][0] or not self.proyectos[pid][1]:
                raise ErrorCatalogo(f"Falta nombre o gestion_actual en {pid}.")

        self.destinos = {}
        ocupados = set()
        hojas_ocupadas = set()
        for fila in registros(tablas, "DESTINOS_PROYECTO", ("proyecto_id", "gestion", "drive_id", "item_id", "hoja", "tabla", "formato_fecha")):
            pid = identificador(fila["proyecto_id"])
            gestion = texto(fila["gestion"])
            if pid not in self.proyectos or not gestion or (pid, gestion) in self.destinos:
                raise ErrorCatalogo("Destino duplicado, gestión vacía o proyecto inexistente.")
            drive, item, tabla = (identificador(fila[c]) for c in ("drive_id", "item_id", "tabla"))
            hoja, formato = texto(fila["hoja"]), texto(fila["formato_fecha"]).upper()
            if not hoja or formato not in ("PE", "US"):
                raise ErrorCatalogo("Cada destino necesita hoja y formato_fecha PE o US.")
            clave = (drive, item, tabla.casefold())
            clave_hoja = (drive, item, hoja.casefold())
            if clave in ocupados or clave_hoja in hojas_ocupadas or (drive, item) == archivo_config:
                raise ErrorCatalogo("Cada gestión necesita su propia hoja y tabla; no uses el catálogo como destino contable.")
            ocupados.add(clave)
            hojas_ocupadas.add(clave_hoja)
            self.destinos[pid, gestion] = Destino(pid, self.proyectos[pid][0], gestion, drive, item, hoja, tabla, formato)
        # Evitar que dos proyectos compartan accidentalmente el mismo archivo.
        archivos = {}
        archivos_por_proyecto = {}
        for destino in self.destinos.values():
            clave = destino.drive_id, destino.item_id
            if clave in archivos and archivos[clave] != destino.proyecto_id:
                raise ErrorCatalogo("Cada proyecto debe tener su propio archivo Excel.")
            anterior = archivos_por_proyecto.get(destino.proyecto_id)
            if anterior is not None and anterior != clave:
                raise ErrorCatalogo("Las gestiones de un proyecto deben ser hojas del mismo archivo Excel.")
            archivos[clave] = destino.proyecto_id
            archivos_por_proyecto[destino.proyecto_id] = clave
        for pid, (_, gestion, habilitado) in self.proyectos.items():
            if habilitado and (pid, gestion) not in self.destinos:
                raise ErrorCatalogo(f"No existe destino para la gestión actual de {pid}.")

        self.permisos = {}
        for fila in registros(tablas, "ACCESOS_PROYECTO", ("telegram_id", "proyecto_id", "activo")):
            uid, pid = telegram_id(fila["telegram_id"]), identificador(fila["proyecto_id"])
            if uid not in self.usuarios or pid not in self.proyectos:
                raise ErrorCatalogo("Un permiso apunta a una persona o proyecto inexistente.")
            clave = uid, pid
            self.permisos[clave] = self.permisos.get(clave, True) and activo(fila["activo"])

    def destinos_usuario(self, uid):
        uid = telegram_id(uid)
        if not self.usuarios.get(uid, False):
            return []
        return sorted((self.destinos[pid, gestion]
                       for pid, (_, gestion, habilitado) in self.proyectos.items()
                       if habilitado and self.permisos.get((uid, pid), False)),
                      key=lambda destino: destino.proyecto_id)

    def validar_destino(self, uid, elegido):
        if elegido not in self.destinos_usuario(uid):
            raise ErrorCatalogo("El proyecto, gestión, destino o permiso cambió. Inicia una operación nueva.")
        return elegido
