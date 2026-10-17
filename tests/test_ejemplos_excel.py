"""Contrato de los ejemplos distribuidos con el bot. Solo lectura local."""
from pathlib import Path
import unittest

from catalogo import Catalogo, leer_tablas_xlsx


RAIZ = Path(__file__).resolve().parents[1]
EJEMPLOS = RAIZ / "ejemplos" / "excel"


class EjemplosExcel(unittest.TestCase):
    def libro(self, nombre, valores=False):
        from openpyxl import load_workbook
        libro = load_workbook(EJEMPLOS / nombre, data_only=valores)
        self.addCleanup(libro.close)
        return libro

    def test_catalogo_valido_sin_habilitar_personas_o_proyectos_ficticios(self):
        tablas = leer_tablas_xlsx((EJEMPLOS / "Accesos_ejemplo.xlsx").read_bytes())
        self.assertEqual(set(tablas), {"USUARIOS_AUTORIZADOS", "PROYECTOS", "DESTINOS_PROYECTO", "ACCESOS_PROYECTO"})
        catalogo = Catalogo(tablas)
        self.assertFalse(any(catalogo.usuarios.values()))
        self.assertFalse(any(proyecto[2] for proyecto in catalogo.proyectos.values()))
        self.assertEqual(catalogo.destinos_usuario("123456789"), [])
        self.assertTrue(all(destino.formato_fecha == "PE" for destino in catalogo.destinos.values()))

    def test_catalogo_identifica_hojas_y_tablas_del_archivo_contable(self):
        catalogo = Catalogo(leer_tablas_xlsx((EJEMPLOS / "Accesos_ejemplo.xlsx").read_bytes()))
        libro = self.libro("Proyecto_X_ejemplo.xlsx")
        self.assertEqual(libro.sheetnames, ["2025", "2026"])
        for gestion in libro.sheetnames:
            destino = catalogo.destinos["proyecto_x", gestion]
            self.assertIn(destino.tabla, libro[destino.hoja].tables)
            tabla = libro[destino.hoja].tables[destino.tabla]
            self.assertEqual([col.name for col in tabla.tableColumns], [
                "CODIGO", "FECHA", "NRO_OPERACION", "CONCEPTO", "TIPO", "ING_EG",
                "MOTIVO", "ACREEDOR", "DEUDOR", "ESTADO", "INGRESO", "EGRESO", "SALDO"])

    def test_formulas_y_saldos_de_los_ejemplos_se_conservan_en_archivo(self):
        libro = self.libro("Proyecto_X_ejemplo.xlsx")
        calculado = self.libro("Proyecto_X_ejemplo.xlsx", valores=True)
        for gestion in libro.sheetnames:
            hoja = libro[gestion]
            self.assertEqual(hoja["B3"].value, 1000)
            saldo = hoja["B3"].value
            for fila in (7, 8, 9):
                saldo += hoja[f"K{fila}"].value - hoja[f"L{fila}"].value
                self.assertAlmostEqual(calculado[gestion][f"M{fila}"].value, saldo)
                fecha = hoja[f"B{fila}"].value
                numero = hoja[f"C{fila}"].value
                self.assertEqual(calculado[gestion][f"A{fila}"].value, f"O{fecha:%d%m%y}{numero:02d}")
            for fila in (7, 8, 9, 10):
                self.assertEqual(hoja[f"A{fila}"].data_type, "f")
                self.assertEqual(hoja[f"M{fila}"].data_type, "f")
            self.assertIn(calculado[gestion]["M10"].value, (None, ""))
            self.assertIn(calculado[gestion]["A10"].value, (None, ""))


if __name__ == "__main__":
    unittest.main()
